import openpyxl
import json
import os

def to_float(val):
    if val is None:
        return None
    try:
        return float(str(val).strip())
    except (ValueError, TypeError):
        return None

def to_int(val, default=3):
    if val is None:
        return default
    try:
        return int(float(str(val).strip()))
    except (ValueError, TypeError):
        return default

def normalize_course_type(ctype_str):
    if not ctype_str:
        return "THEORY"
    c = str(ctype_str).strip().upper()
    if "LAB" in c or "PRACTICAL" in c:
        return "LAB"
    if "PROJ" in c or "MINI" in c:
        return "PROJECT"
    return "THEORY"

CANONICAL_CODE_MAP = {
    'ANH-MPCECDCM701T': 'N-PCCCM701T',
    'UNT-APMCCCM701T': 'N-PCCCM701T',
    'ANH-MPCECDCM702T': 'N-PCCCM702T',
    'UNT-APMCCCM702T': 'N-PCCCM702T',
    'ANH-MPCECDCM702P': 'N-PCCCM702P',
    'UNT-APMCCCM702P': 'N-PCCCM702P',
    'ANH-MPEECDCM701T': 'N-PECCM701T',
    'ANH-MPEECDCM701P': 'N-PECCM701P',
    'ANH-MPEECDCM704T': 'N-PECCM704T',
    'ANH-MPREODJCM701': 'N-PROJCM701',
    'UNT-APMROJCM701': 'N-PROJCM701',
    'UNT-APMECCM703T': 'N-PECCM703T',
    'UNT-APMECCM703P': 'N-PECCM703P',
    'UNT-APMECCM706T': 'N-PECCM706T',
    'ANH-MMEDDMEF701T': 'N-MDMEF701T'
}

CANONICAL_TITLE_MAP = {
    'N-PCCCM701T': 'Neural Networks and Deep Learning',
    'N-PCCCM702T': 'Machine Vision',
    'N-PCCCM702P': 'Machine Vision Lab',
    'N-PECCM701T': 'Introduction To Cyber Security',
    'N-PECCM701P': 'Introduction To Cyber Security Lab',
    'N-PECCM704T': 'Cyber Forensics and Investigation',
    'N-PROJCM701': 'Project-I',
    'N-PECCM703T': 'Pattern Recognition',
    'N-PECCM703P': 'Pattern Recognition Lab',
    'N-PECCM706T': 'Explainable AI',
    'N-MDMEF701T': 'Investment Analysis & Portfolio Management',
    'N-MDMIT701P': 'Mini Project in IoT Product Development & Testing'
}

def parse_excel():
    excel_candidates = [
        "../Merged_Student_Subject_Registration_2026-27_With_Faculty.xlsx",
        "Merged_Student_Subject_Registration_2026-27_With_Faculty.xlsx",
        "../Merged_Student_Subject_Registration_2026-27.xlsx",
        "Merged_Student_Subject_Registration_2026-27.xlsx"
    ]
    excel_path = None
    for p in excel_candidates:
        if os.path.exists(p):
            excel_path = p
            break
            
    if not excel_path:
        raise FileNotFoundError("Could not find student subject registration Excel file.")
        
    print(f"Loading Excel file from {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb["Merged Complete Data"]
    
    students = {}
    courses = {}
    faculties = {}
    student_courses = []
    
    # Iterate from row 2 to sheet.max_row
    for r in range(2, sheet.max_row + 1):
        prn = sheet.cell(row=r, column=3).value
        email = sheet.cell(row=r, column=11).value
        
        # Skip empty rows
        if not prn or not email:
            continue
            
        prn = str(prn).strip()
        email = str(email).strip().lower()
        
        full_name = sheet.cell(row=r, column=10).value or sheet.cell(row=r, column=4).value
        full_name = str(full_name).strip() if full_name else "Unknown Student"
        
        raw_code = sheet.cell(row=r, column=5).value
        raw_name = sheet.cell(row=r, column=6).value
        raw_type = sheet.cell(row=r, column=7).value
        raw_credits = sheet.cell(row=r, column=8).value
        
        tenth = to_float(sheet.cell(row=r, column=12).value)
        twelfth = to_float(sheet.cell(row=r, column=13).value)
        sem1 = to_float(sheet.cell(row=r, column=14).value)
        sem2 = to_float(sheet.cell(row=r, column=15).value)
        sem3 = to_float(sheet.cell(row=r, column=16).value)
        sem4 = to_float(sheet.cell(row=r, column=17).value)
        sem5 = to_float(sheet.cell(row=r, column=18).value)
        sem6 = to_float(sheet.cell(row=r, column=19).value)
        
        backlog_val = sheet.cell(row=r, column=20).value
        backlogs = str(backlog_val).strip() if backlog_val and str(backlog_val).strip().upper() not in ["NONE", "N/A", "NIL", "NO", "0"] else None
        
        internship_val = sheet.cell(row=r, column=21).value
        internships = str(internship_val).strip() if internship_val and str(internship_val).strip().upper() not in ["NONE", "N/A", "NIL", "NO", "0"] else None

        faculty_name_val = sheet.cell(row=r, column=22).value if sheet.max_column >= 22 else None
        faculty_email_val = sheet.cell(row=r, column=23).value if sheet.max_column >= 23 else None

        faculty_name = str(faculty_name_val).strip() if faculty_name_val else "Prof. Sarah Jenkins"
        faculty_email = str(faculty_email_val).strip().lower() if faculty_email_val else "faculty1@sbjit.edu.in"
        
        if not faculty_name.startswith("Prof."):
            faculty_name = f"Prof. {faculty_name}"

        # Register faculty
        if faculty_email not in faculties:
            faculties[faculty_email] = {
                "fullName": faculty_name,
                "email": faculty_email,
                "department": "Artificial Intelligence & Machine Learning",
            }

        # Canonical normalization
        clean_raw_code = str(raw_code).strip() if raw_code else ""
        canonical_code = CANONICAL_CODE_MAP.get(clean_raw_code, clean_raw_code)
        course_title = CANONICAL_TITLE_MAP.get(canonical_code, str(raw_name).strip() if raw_name else "Curriculum Course")
        course_type = normalize_course_type(raw_type)
        credits_val = to_int(raw_credits, default=1 if course_type == "LAB" else 4 if course_type == "PROJECT" else 3)

        # Add unique student
        if email not in students:
            students[email] = {
                "prn": prn,
                "fullName": full_name,
                "email": email,
                "tenthPercentage": tenth,
                "twelfthPercentage": twelfth,
                "sem1Cgpa": sem1,
                "sem2Cgpa": sem2,
                "sem3Cgpa": sem3,
                "sem4Cgpa": sem4,
                "sem5Cgpa": sem5,
                "sem6Cgpa": sem6,
                "backlogs": backlogs,
                "internships": internships,
                "tgMentorName": "Prof. Bhushan Manjrekar" if int(prn.replace('CM23', '') if 'CM23' in prn and prn.replace('CM23', '').isdigit() else '1') <= 15 else "Prof. Sujata Sardare" if int(prn.replace('CM23', '') if 'CM23' in prn and prn.replace('CM23', '').isdigit() else '1') <= 30 else "Prof. Nikhil Sakhare" if int(prn.replace('CM23', '') if 'CM23' in prn and prn.replace('CM23', '').isdigit() else '1') <= 45 else "Prof. Shweta Bokade"
            }
            
        # Add unique canonical course
        if canonical_code:
            if canonical_code not in courses:
                courses[canonical_code] = {
                    "courseCode": canonical_code,
                    "title": course_title,
                    "courseType": course_type,
                    "credits": credits_val,
                    "facultyName": faculty_name,
                    "facultyEmail": faculty_email
                }
            
            # Map course to student
            student_courses.append({
                "prn": prn,
                "courseCode": canonical_code,
                "courseName": course_title,
                "courseType": course_type,
                "credits": credits_val,
                "facultyEmail": faculty_email
            })
            
    # Export to JSON
    output_dir = "storage"
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
        
    output_path = os.path.join(output_dir, "roster_data.json")
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump({
            "faculties": list(faculties.values()),
            "students": list(students.values()),
            "courses": list(courses.values()),
            "studentCourses": student_courses
        }, f, indent=2)
        
    print(f"Successfully processed {len(faculties)} faculty members, {len(students)} students, and {len(courses)} unified canonical courses.")
    print(f"Unified Courses: {[c['courseCode'] + ' (' + c['title'] + ' - ' + c['facultyName'] + ')' for c in courses.values()]}")

if __name__ == "__main__":
    parse_excel()
